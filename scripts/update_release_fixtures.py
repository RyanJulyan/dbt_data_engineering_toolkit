#!/usr/bin/env python3
"""Refresh release fixtures that must stay in lockstep with version bumps."""

from __future__ import annotations

import os
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

from dbt_data_engineering_toolkit_compiler.version import COMPILER_VERSION


ROOT = Path(__file__).resolve().parents[1]
PYTHON_SOURCE = ROOT / "python" / "src"
WORKBOOKS = (
    ROOT / "python/src/dbt_data_engineering_toolkit_compiler/resources/data_product_sample.xlsx",
    ROOT
    / "python/src/dbt_data_engineering_toolkit_compiler/resources/data_product_multi_source_sample.xlsx",
)
RELEASE_OUTPUTS = [
    *WORKBOOKS,
    ROOT / "examples/compiler/customer_360",
    ROOT / "examples/compiler/customer_accounts",
]


def update_workbook_template_version(path: Path) -> None:
    workbook = load_workbook(path)
    try:
        if "_DET Metadata" not in workbook.sheetnames:
            raise ValueError(f"Workbook is missing _DET Metadata: {path}")
        workbook["_DET Metadata"]["B4"] = COMPILER_VERSION
        workbook.save(path)
    finally:
        workbook.close()


def run_command(args: list[str]) -> None:
    env = dict(os.environ)
    pythonpath = env.get("PYTHONPATH")
    env["PYTHONPATH"] = (
        f"{PYTHON_SOURCE}{os.pathsep}{pythonpath}"
        if pythonpath
        else str(PYTHON_SOURCE)
    )
    result = subprocess.run(
        args,
        cwd=ROOT,
        env=env,
        capture_output=True,
        text=True,
        check=False,
    )
    if result.returncode != 0:
        sys.stdout.write(result.stdout)
        sys.stderr.write(result.stderr)
        raise SystemExit(result.returncode)


def stage_release_outputs() -> None:
    run_command(["git", "add", "--", *[str(path.relative_to(ROOT)) for path in RELEASE_OUTPUTS]])


def main() -> int:
    for workbook_path in WORKBOOKS:
        update_workbook_template_version(workbook_path)
    run_command(
        [
            sys.executable,
            "-m",
            "dbt_data_engineering_toolkit_compiler",
            "workbook",
            "refresh",
            "python/src/dbt_data_engineering_toolkit_compiler/resources/data_product_sample.xlsx",
        ]
    )
    run_command(
        [
            sys.executable,
            "-m",
            "dbt_data_engineering_toolkit_compiler",
            "workbook",
            "refresh",
            "python/src/dbt_data_engineering_toolkit_compiler/resources/data_product_multi_source_sample.xlsx",
        ]
    )
    run_command(
        [
            sys.executable,
            "-m",
            "dbt_data_engineering_toolkit_compiler",
            "generate",
            "python/src/dbt_data_engineering_toolkit_compiler/resources/data_product_sample.xlsx",
            "--project-dir",
            "examples/compiler/customer_360",
            "--prune",
            "--force",
        ]
    )
    run_command(
        [
            sys.executable,
            "-m",
            "dbt_data_engineering_toolkit_compiler",
            "generate",
            "python/src/dbt_data_engineering_toolkit_compiler/resources/data_product_multi_source_sample.xlsx",
            "--project-dir",
            "examples/compiler/customer_accounts",
            "--prune",
            "--force",
        ]
    )
    stage_release_outputs()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())