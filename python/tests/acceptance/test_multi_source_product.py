"""Realistic multi-source product acceptance fixture."""

from __future__ import annotations

from importlib.resources import files
from pathlib import Path

from openpyxl import load_workbook

from dbt_data_engineering_toolkit_compiler.services.emissions.service import EmissionService
from dbt_data_engineering_toolkit_compiler.services.validation.service import (
    SpecificationValidationService,
)
from dbt_data_engineering_toolkit_compiler.services.workbooks.service import (
    WorkbookInterpretationService,
)


def fixture_workbook() -> Path:
    return Path(
        str(
            files("dbt_data_engineering_toolkit_compiler")
            / "resources"
            / "data_product_multi_source_sample.xlsx"
        )
    )


def load_specification(path: Path):
    return WorkbookInterpretationService().load(path)


def validate_specification(specification) -> None:
    SpecificationValidationService().validate(specification)


def emit_all(specification):
    return EmissionService().emit_all(specification)


def test_multi_source_workbook_compiles_joined_readable_sql() -> None:
    spec = load_specification(fixture_workbook())
    validate_specification(spec)
    assert len(spec.sources) == 2
    assert len(spec.models) == 2
    assert len(spec.relationships) == 1
    assert len(spec.lookups) == 2

    artifacts = {item.path.as_posix(): item.content for item in emit_all(spec)}
    staging = artifacts["models/staging/stg_customers.sql"]
    mart = artifacts["models/marts/customer_account_mart.sql"]
    sources = artifacts["models/sources.yml"]
    assert "{{ de_toolkit.mapping(" in staging
    assert "{{ de_toolkit.string_title('customer_name') }}" in staging
    assert "transformed_02 as" in staging
    assert "from source__stg_customers as stg_customers" in mart
    assert "left join source__raw_accounts as raw_accounts" in mart
    assert "from joined" not in mart
    assert "- unique" in sources


def test_workbook_has_contextual_lists_and_release_metadata() -> None:
    workbook = load_workbook(fixture_workbook(), data_only=False)
    try:
        assert workbook["_DET Metadata"]["B4"].value == "2.3.0"
        assert len(str(workbook["_DET Metadata"]["B7"].value)) == 64
        assert workbook["_DET Metadata"]["B8"].value == "3.1.0"
        for name in ("_DET Lists", "_DET Metadata", "_DET Context", "_DET Raw ODCS"):
            assert workbook[name].sheet_state == "veryHidden"
            assert workbook[name].protection.sheet
        assert workbook["Quality"].sheet_state == "visible"
        assert {
            name
            for name in workbook.sheetnames
            if "quality" in name.casefold() or name.startswith("Operational ")
        } == {
            "Quality",
            "Operational Validation",
            "Operational Parameters",
        }
        assert workbook["Quality"]["N4"].value == "Rule ID"
        assert workbook["Quality"]["P4"].value == "ODCS Passthrough (JSON)"
        assert "_DET_MODEL_customer_account_mart" in workbook.defined_names
        assert "_DET_INPUTS_customer_account_mart" in workbook.defined_names
        assert "_DET_REL_raw_accounts" in workbook.defined_names
        mapping_formulas = [
            item.formula1 for item in workbook["DET Mapping"].data_validations.dataValidation
        ]
        relationship_formulas = [
            item.formula1 for item in workbook["DET Relationships"].data_validations.dataValidation
        ]
        quality_formulas = [
            item.formula1 for item in workbook["Quality"].data_validations.dataValidation
        ]
        assert any("_DET_MODEL_" in formula for formula in mapping_formulas)
        assert any("_DET_INPUTS_" in formula for formula in relationship_formulas)
        assert any("_DET_REL_" in formula for formula in relationship_formulas)
        assert any("_DET_MODEL_" in formula for formula in quality_formulas)
        schema = workbook["Schema customer_account_mart"]
        assert not schema.column_dimensions["N"].hidden
        assert schema.column_dimensions["O"].hidden
        assert not schema.column_dimensions["AM"].hidden
    finally:
        workbook.close()
