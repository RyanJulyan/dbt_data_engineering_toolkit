"""Workbook broker that isolates openpyxl from DET workbook interpretation."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Iterator, Protocol, TypeAlias

from openpyxl import load_workbook

from ..errors import InvalidWorkbookError

CellValue: TypeAlias = str | int | float | bool | date | datetime | None


def _cell_value(value: object) -> CellValue:
    """Normalize library-specific cell types before they cross the broker boundary."""

    if value is None or isinstance(value, (str, int, float, bool, date, datetime)):
        return value
    return str(value)


@dataclass(frozen=True, slots=True)
class SheetData:
    """Technology-neutral, one-indexed snapshot of a worksheet."""

    name: str
    values: tuple[tuple[CellValue, ...], ...]

    @property
    def max_row(self) -> int:
        return len(self.values)

    @property
    def max_column(self) -> int:
        return max((len(row) for row in self.values), default=0)

    def cell(self, row: int, column: int) -> CellValue:
        if row < 1 or column < 1 or row > self.max_row:
            return None
        values = self.values[row - 1]
        return values[column - 1] if column <= len(values) else None

    def row(self, row: int) -> tuple[CellValue, ...]:
        return self.values[row - 1] if 1 <= row <= self.max_row else ()

    def iter_rows(self, *, min_row: int = 1) -> Iterator[tuple[CellValue, ...]]:
        yield from self.values[max(min_row - 1, 0) :]


@dataclass(frozen=True, slots=True)
class WorkbookDocument:
    """Read-only workbook snapshot consumed by interpretation services."""

    sheets: dict[str, SheetData]

    @property
    def sheetnames(self) -> tuple[str, ...]:
        return tuple(self.sheets)

    def sheet(self, name: str) -> SheetData:
        return self.sheets[name]


class WorkbookBroker(Protocol):
    def read(self, path: Path) -> WorkbookDocument: ...


class OpenpyxlWorkbookBroker:
    """Production workbook reader; openpyxl never leaks above this class."""

    def read(self, path: Path) -> WorkbookDocument:
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:  # openpyxl exposes several format-specific exceptions
            raise InvalidWorkbookError(f"Could not open {path.name}: {exc}") from exc
        try:
            sheets = {
                name: SheetData(
                    name=name,
                    values=tuple(
                        tuple(_cell_value(cell.value) for cell in row)
                        for row in workbook[name].iter_rows()
                    ),
                )
                for name in workbook.sheetnames
            }
            return WorkbookDocument(sheets=sheets)
        finally:
            workbook.close()
